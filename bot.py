#!/usr/bin/env python3
"""VOLP assignment tracker + Telegram submission bot — multi-user edition."""

import argparse
import asyncio
import hashlib
import os
import sqlite3
import sys
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Dict, List, Optional
from urllib.parse import quote

import requests
from dotenv import load_dotenv
from playwright.async_api import async_playwright
from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    ApplicationBuilder,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    ConversationHandler,
    MessageHandler,
    filters,
)

# Assignment type hrefs that exist in VOLP course content pages
ASSIGNMENT_TYPE_HREFS = [
    "/learner-subjective-assignment",
    "/learner-handson-assignment",
    "/learner-cie-assignment",
    "/learner-mcq-assignment",
    "/learner-matchthepair-assignment",
    "/learner-singleword-assignment",
]

# Conversation states for /start registration flow
WAIT_USERNAME = 1
WAIT_PASSWORD = 2

# Free-tier models on OpenRouter — tried in order, first 200-OK wins
OPENROUTER_FREE_MODELS = [
    "google/gemma-3-27b-it:free",
    "stepfun/step-3.5-flash:free",
    "arcee-ai/trinity-mini:free",
    "nvidia/nemotron-nano-9b-v2:free",
]
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------


@dataclass
class Assignment:
    entity_key: str
    title: str
    course: str
    due_raw: str
    due_iso: str
    source_url: str      # SPA URL, e.g. /learner-subjective-assignment
    question_no: str
    overview_url: str    # course overview URL for re-navigation
    submitted_file: str = ""   # filename if the user has uploaded something
    full_text: str = ""        # complete problem statement (untruncated)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# State store — multi-user edition
# ---------------------------------------------------------------------------


class StateStore:
    def __init__(self, db_path: str):
        db_dir = os.path.dirname(os.path.abspath(db_path))
        os.makedirs(db_dir, exist_ok=True)
        self.conn = sqlite3.connect(db_path, check_same_thread=False)
        self.conn.row_factory = sqlite3.Row
        self._init_schema()

    def _init_schema(self) -> None:
        # Users table
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
              chat_id        TEXT PRIMARY KEY,
              volp_username  TEXT NOT NULL,
              volp_password  TEXT NOT NULL,
              openrouter_key TEXT DEFAULT '',
              status         TEXT NOT NULL DEFAULT 'pending',
              registered_at  TEXT NOT NULL,
              last_scan_at   TEXT DEFAULT ''
            )
            """
        )
        # Meta table
        self.conn.execute(
            """
            CREATE TABLE IF NOT EXISTS meta (
              key   TEXT PRIMARY KEY,
              value TEXT NOT NULL
            )
            """
        )
        # Assignments table — create fresh or migrate from single-user schema
        self._init_assignments_table()
        self.conn.commit()

    def _init_assignments_table(self) -> None:
        """Create or migrate the assignments table for multi-user support."""
        tables = {
            row[0]
            for row in self.conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table'"
            ).fetchall()
        }

        if "assignments" not in tables:
            # Fresh install — create with composite PK
            self.conn.execute(
                """
                CREATE TABLE assignments (
                  entity_key     TEXT NOT NULL,
                  user_id        TEXT NOT NULL DEFAULT '',
                  title          TEXT NOT NULL,
                  course         TEXT NOT NULL,
                  due_raw        TEXT NOT NULL,
                  due_iso        TEXT,
                  source_url     TEXT,
                  overview_url   TEXT,
                  question_no    TEXT,
                  submitted_file TEXT DEFAULT '',
                  full_text      TEXT DEFAULT '',
                  first_seen_at  TEXT NOT NULL,
                  last_seen_at   TEXT NOT NULL,
                  PRIMARY KEY (entity_key, user_id)
                )
                """
            )
            return

        # Check if already on multi-user schema (user_id column present)
        existing_cols = {
            row[1]
            for row in self.conn.execute("PRAGMA table_info(assignments)").fetchall()
        }
        if "user_id" in existing_cols:
            return  # Already migrated

        # Migrate: recreate table with composite PK, preserving existing data as user_id=''
        print("[db] migrating assignments table to multi-user schema...")
        self.conn.executescript(
            """
            BEGIN;

            ALTER TABLE assignments RENAME TO assignments_v1;

            CREATE TABLE assignments (
              entity_key     TEXT NOT NULL,
              user_id        TEXT NOT NULL DEFAULT '',
              title          TEXT NOT NULL,
              course         TEXT NOT NULL,
              due_raw        TEXT NOT NULL,
              due_iso        TEXT,
              source_url     TEXT,
              overview_url   TEXT,
              question_no    TEXT,
              submitted_file TEXT DEFAULT '',
              full_text      TEXT DEFAULT '',
              first_seen_at  TEXT NOT NULL,
              last_seen_at   TEXT NOT NULL,
              PRIMARY KEY (entity_key, user_id)
            );

            INSERT INTO assignments
              SELECT
                entity_key, '',
                title, course, due_raw,
                COALESCE(due_iso, ''),
                COALESCE(source_url, ''),
                COALESCE(overview_url, ''),
                COALESCE(question_no, ''),
                COALESCE(submitted_file, ''),
                COALESCE(full_text, ''),
                first_seen_at, last_seen_at
              FROM assignments_v1;

            DROP TABLE assignments_v1;

            COMMIT;
            """
        )
        print("[db] migration complete")

    # --- User methods ---

    def register_user(
        self, chat_id: str, username: str, password: str, status: str = "pending"
    ) -> None:
        self.conn.execute(
            """
            INSERT INTO users(chat_id, volp_username, volp_password, status, registered_at)
            VALUES(?,?,?,?,?)
            ON CONFLICT(chat_id) DO UPDATE SET
              volp_username=excluded.volp_username,
              volp_password=excluded.volp_password,
              status=excluded.status
            """,
            (chat_id, username, password, status, now_iso()),
        )
        self.conn.commit()

    def get_user(self, chat_id: str) -> Optional[sqlite3.Row]:
        return self.conn.execute(
            "SELECT * FROM users WHERE chat_id = ?", (chat_id,)
        ).fetchone()

    def get_all_users(self, status: Optional[str] = None) -> List[sqlite3.Row]:
        if status:
            return self.conn.execute(
                "SELECT * FROM users WHERE status = ?", (status,)
            ).fetchall()
        return self.conn.execute("SELECT * FROM users").fetchall()

    def update_user_status(self, chat_id: str, status: str) -> None:
        self.conn.execute(
            "UPDATE users SET status = ? WHERE chat_id = ?", (status, chat_id)
        )
        self.conn.commit()

    def update_user_last_scan(self, chat_id: str) -> None:
        self.conn.execute(
            "UPDATE users SET last_scan_at = ? WHERE chat_id = ?", (now_iso(), chat_id)
        )
        self.conn.commit()

    def set_user_key(self, chat_id: str, key: str) -> None:
        self.conn.execute(
            "UPDATE users SET openrouter_key = ? WHERE chat_id = ?", (key, chat_id)
        )
        self.conn.commit()

    def migrate_legacy_user(self, chat_id: str) -> None:
        """Assign orphan assignments (user_id='') to this user."""
        self.conn.execute(
            "UPDATE assignments SET user_id = ? WHERE user_id = ''", (chat_id,)
        )
        self.conn.commit()

    # --- Meta methods ---

    def get_meta(self, key: str) -> str:
        row = self.conn.execute("SELECT value FROM meta WHERE key = ?", (key,)).fetchone()
        return row["value"] if row else ""

    def set_meta(self, key: str, value: str) -> None:
        self.conn.execute(
            "INSERT INTO meta(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, value),
        )
        self.conn.commit()

    # --- Assignment methods ---

    def load_assignments(self, user_id: str) -> Dict[str, sqlite3.Row]:
        rows = self.conn.execute(
            "SELECT * FROM assignments WHERE user_id = ?", (user_id,)
        ).fetchall()
        return {row["entity_key"]: row for row in rows}

    def get_by_entity_key(self, entity_key: str, user_id: str) -> Optional[sqlite3.Row]:
        return self.conn.execute(
            "SELECT * FROM assignments WHERE entity_key = ? AND user_id = ?",
            (entity_key, user_id),
        ).fetchone()

    def upsert_assignments(self, assignments: List[Assignment], user_id: str) -> None:
        ts = now_iso()
        for item in assignments:
            self.conn.execute(
                """
                INSERT INTO assignments(entity_key, user_id, title, course, due_raw, due_iso,
                    source_url, overview_url, question_no, submitted_file, full_text,
                    first_seen_at, last_seen_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(entity_key, user_id) DO UPDATE SET
                  title=excluded.title, course=excluded.course,
                  due_raw=excluded.due_raw, due_iso=excluded.due_iso,
                  source_url=excluded.source_url, overview_url=excluded.overview_url,
                  question_no=excluded.question_no,
                  submitted_file=excluded.submitted_file,
                  full_text=excluded.full_text,
                  last_seen_at=excluded.last_seen_at
                """,
                (
                    item.entity_key, user_id, item.title, item.course,
                    item.due_raw, item.due_iso, item.source_url,
                    item.overview_url, item.question_no, item.submitted_file,
                    item.full_text, ts, ts,
                ),
            )
        self.conn.commit()


# ---------------------------------------------------------------------------
# Simple one-shot Telegram sender (--once mode only)
# ---------------------------------------------------------------------------


class TelegramClient:
    def __init__(self, token: str, chat_id: str):
        self.token = token
        self.chat_id = chat_id

    def enabled(self) -> bool:
        return bool(self.token and self.chat_id)

    def send(self, text: str) -> bool:
        if not self.enabled():
            return False
        try:
            r = requests.post(
                f"https://api.telegram.org/bot{self.token}/sendMessage",
                json={"chat_id": self.chat_id, "text": text, "disable_web_page_preview": True},
                timeout=20,
            )
            r.raise_for_status()
            return bool(r.json().get("ok"))
        except Exception as exc:
            print(f"[telegram] send failed: {exc}")
            return False


# ---------------------------------------------------------------------------
# Shared Playwright login helper
# ---------------------------------------------------------------------------


async def _playwright_login(page, login_url: str, username: str, password: str) -> None:
    await page.goto(login_url, wait_until="domcontentloaded", timeout=30000)
    await page.wait_for_load_state("networkidle", timeout=30000)

    if not await page.locator('input[type="password"]').count() and "login" not in page.url.lower():
        return  # already logged in

    if not username or not password:
        raise RuntimeError("Login required but credentials not provided")

    print("[auth] logging in...")
    for sel in [
        'input[placeholder*="email or login" i]',
        'input[name="username"]', 'input[name="email"]',
        'input[type="email"]', 'input[type="text"]',
    ]:
        if await page.locator(sel).count() > 0:
            await page.fill(sel, username)
            break

    await page.fill('input[type="password"]', password)

    for sel in [
        'button.btn-sign-in', 'button:has-text("SIGN IN")',
        'button[type="submit"]', 'button:has-text("Login")',
    ]:
        if await page.locator(sel).count() > 0:
            await page.click(sel)
            break

    await page.wait_for_load_state("networkidle", timeout=30000)
    await page.wait_for_timeout(2500)

    if await page.locator('input[type="password"]').count() or "login" in page.url.lower():
        raise RuntimeError("Login failed — check credentials")
    print("[auth] logged in")


async def _verify_volp_credentials(username: str, password: str, config: dict) -> bool:
    """Verify VOLP credentials by attempting a headless login. Returns True if valid."""
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=config.get("headless", True))
        ctx = await browser.new_context()
        page = await ctx.new_page()
        try:
            await _playwright_login(page, config["login_url"], username, password)
            return True
        except RuntimeError:
            return False
        finally:
            await browser.close()


# ---------------------------------------------------------------------------
# VOLP Scanner
# ---------------------------------------------------------------------------


class VolpScanner:
    def __init__(
        self,
        base_url: str,
        login_url: str,
        start_url: str,
        username: str,
        password: str,
        headless: bool,
        max_courses: int,
        tz_offset_minutes: int,
    ):
        self.base_url = base_url.rstrip("/")
        self.login_url = login_url
        self.start_url = start_url
        self.username = username
        self.password = password
        self.headless = headless
        self.max_courses = max_courses
        self.tz_offset_minutes = tz_offset_minutes

    async def scan(self) -> List[Assignment]:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=self.headless)
            ctx = await browser.new_context()
            await ctx.route("**/*", _route_filter)
            page = await ctx.new_page()

            await _playwright_login(page, self.login_url, self.username, self.password)
            await page.goto(self.start_url, wait_until="networkidle", timeout=30000)
            await page.wait_for_timeout(3000)

            btn_count = await page.locator('button:has-text("VIEW COURSE")').count()
            if btn_count == 0:
                await page.reload(wait_until="networkidle", timeout=30000)
                await page.wait_for_timeout(3000)
                btn_count = await page.locator('button:has-text("VIEW COURSE")').count()

            total = min(btn_count, self.max_courses)
            print(f"[scan] {btn_count} courses found, scanning {total}")

            all_assignments: List[Assignment] = []
            for index in range(total):
                try:
                    assignments = await self._scan_course(page, index)
                    all_assignments.extend(assignments)
                    print(f"[scan] course {index + 1}/{total}: {len(assignments)} assignments")
                except Exception as exc:
                    print(f"[scan] course {index + 1} failed: {exc}")

            await browser.close()
            return _dedupe(all_assignments)

    async def _scan_course(self, page, index: int) -> List[Assignment]:
        await page.goto(self.start_url, wait_until="networkidle", timeout=30000)
        await page.wait_for_timeout(2500)

        for _ in range(6):
            if await page.locator('button:has-text("VIEW COURSE")').count() > 0:
                break
            await page.wait_for_timeout(1500)

        btn = page.locator('button:has-text("VIEW COURSE")').nth(index)
        if await btn.count() == 0:
            return []
        await btn.click()
        await page.wait_for_load_state("networkidle", timeout=30000)
        await page.wait_for_timeout(2500)
        overview_url = page.url

        print(f"[scan]   overview: {overview_url}")

        await page.wait_for_selector('.v-tab, [role=tab]', timeout=10000)
        if not await _click_tab(page, "COURSE CONTENTS"):
            print(f"[scan]   no COURSE CONTENTS tab")
            return []
        await page.wait_for_timeout(2500)

        hrefs = await page.evaluate(
            """
            (hrefs) => Array.from(document.querySelectorAll('a[href]'))
                .map(el => el.getAttribute('href'))
                .filter(h => hrefs.includes(h))
            """,
            ASSIGNMENT_TYPE_HREFS,
        )
        hrefs = list(dict.fromkeys(hrefs))
        print(f"[scan]   assignment hrefs: {hrefs}")

        results: List[Assignment] = []
        for i, href in enumerate(hrefs):
            try:
                if i > 0:
                    await page.go_back()
                    await page.wait_for_load_state("networkidle", timeout=30000)
                    await page.wait_for_timeout(2000)

                    if await page.locator(f'a[href="{href}"]').count() == 0:
                        await page.goto(overview_url, wait_until="networkidle", timeout=30000)
                        await page.wait_for_timeout(2000)
                        try:
                            await page.wait_for_selector('.v-tab', timeout=8000)
                        except Exception:
                            pass
                        await _click_tab(page, "COURSE CONTENTS")
                        await page.wait_for_timeout(2000)

                link = page.locator(f'a[href="{href}"]').first
                if await link.count() == 0:
                    print(f"[scan] link {href} not found, skipping")
                    continue
                await link.click(timeout=5000)
                await page.wait_for_load_state("networkidle", timeout=30000)
                await page.wait_for_timeout(2500)

                print(f"[scan]   extracting {page.url}")
                extracted = await self._extract_assignments(page, overview_url)
                print(f"[scan]   {href}: {len(extracted)} assignments")
                results.extend(extracted)
            except Exception as exc:
                print(f"[scan] {href} failed: {exc}")

        return results

    async def _extract_assignments(self, page, overview_url: str) -> List[Assignment]:
        source_url = page.url
        raw = await page.evaluate(
            r"""
            () => {
                const courseEl = document.querySelector('.crs strong');
                const courseName = courseEl
                    ? courseEl.textContent.replace(/course\s*:\s*/i, '').trim()
                    : (document.title || 'Unknown Course');

                const cards = Array.from(document.querySelectorAll('.card-border.v-card, .card-border'));

                return cards.map((card, i) => {
                    const questionEl = card.querySelector('.textq strong');
                    const dueDateEl  = card.querySelector('.textd strong');

                    const questionText = questionEl ? questionEl.textContent.trim() : '';
                    const questionNo   = questionText.replace(/Question\s*No\.?\s*/i, '').trim();

                    const dueDateText  = dueDateEl ? dueDateEl.textContent.trim() : '';
                    const dateMatch    = dueDateText.match(
                        /Due\s*Date\s*:\s*(\d{1,2}\/\d{1,2}\/\d{4}(?:\s+\d{1,2}:\d{2}\s*(?:AM|PM)?)?)/i
                    );
                    const due_raw = dateMatch ? dateMatch[1].trim() : '';

                    let fullText = '';
                    const descEl = card.querySelector('.col-sm-7 span, .col-sm-7.col-12 span');
                    if (descEl) {
                        fullText = descEl.textContent.trim();
                    }
                    if (!fullText) {
                        fullText = card.innerText
                            .split('\n')
                            .map(l => l.trim())
                            .filter(l => l.length > 3
                                && !/question no/i.test(l)
                                && !/due date/i.test(l)
                                && !/marks/i.test(l)
                                && !/NOTE/i.test(l)
                                && !/File size/i.test(l)
                                && !/Answer/i.test(l)
                                && !/SUBMIT/i.test(l)
                            ).join('\n') || questionText;
                    }
                    const title = fullText.slice(0, 140).trim();

                    let submittedFile = '';
                    const fileTextEl = card.querySelector('.v-file-input__text');
                    if (fileTextEl) {
                        const t = fileTextEl.textContent.trim();
                        if (t && !/select\s*file/i.test(t)) {
                            submittedFile = t;
                        }
                    }
                    if (!submittedFile) {
                        const m = card.innerText.match(/Select File[\s\n\r]+([^\n\r]+)[\n\r]+OR/);
                        if (m && m[1].trim() && !/select\s*file/i.test(m[1])) {
                            submittedFile = m[1].trim();
                        }
                    }

                    return { title, full_text: fullText, course: courseName, due_raw, question_no: questionNo, submitted_file: submittedFile };
                }).filter(item => item.question_no);
            }
            """
        )

        items: List[Assignment] = []
        for row in raw:
            due_raw = (row.get("due_raw") or "").strip()
            due_iso = normalize_due_to_iso(due_raw, self.tz_offset_minutes) if due_raw else ""
            key_source = "|".join([
                (row.get("course") or "").strip().lower(),
                (row.get("question_no") or "").strip().lower(),
                source_url.lower(),
            ])
            items.append(
                Assignment(
                    entity_key=hashlib.sha1(key_source.encode()).hexdigest(),
                    title=(row.get("title") or "").strip(),
                    course=(row.get("course") or "Unknown Course").strip(),
                    due_raw=due_raw,
                    due_iso=due_iso,
                    source_url=source_url,
                    question_no=(row.get("question_no") or "").strip(),
                    overview_url=overview_url,
                    submitted_file=(row.get("submitted_file") or "").strip(),
                    full_text=(row.get("full_text") or "").strip(),
                )
            )
        return items


# ---------------------------------------------------------------------------
# VOLP Submitter
# ---------------------------------------------------------------------------


class VolpSubmitter:
    def __init__(self, login_url: str, start_url: str, username: str, password: str, headless: bool):
        self.login_url = login_url
        self.start_url = start_url
        self.username = username
        self.password = password
        self.headless = headless

    async def submit_file(
        self, overview_url: str, source_url: str, question_no: str, file_path: str
    ) -> str:
        """Navigate to the assignment page via SPA context and upload file_path."""
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=self.headless)
            ctx = await browser.new_context()
            page = await ctx.new_page()
            try:
                await _playwright_login(page, self.login_url, self.username, self.password)

                await page.goto(overview_url, wait_until="networkidle", timeout=30000)
                await page.wait_for_timeout(1500)

                if not await _click_tab(page, "COURSE CONTENTS"):
                    return "❌ Could not find COURSE CONTENTS tab"
                await page.wait_for_timeout(1500)

                from urllib.parse import urlparse
                parsed = urlparse(source_url)
                href = parsed.path

                link = page.locator(f'a[href="{href}"]').first
                if await link.count() == 0:
                    return f"❌ Assignment link {href} not found in course contents"
                await link.click(timeout=5000)
                await page.wait_for_load_state("networkidle", timeout=30000)
                await page.wait_for_timeout(2000)

                return await self._do_upload(page, question_no, file_path)
            except Exception as exc:
                return f"❌ Submission failed: {exc}"
            finally:
                await browser.close()

    async def _do_upload(self, page, question_no: str, file_path: str) -> str:
        cards = page.locator(".card-border.v-card, .card-border")
        count = await cards.count()
        target_card = None

        for i in range(count):
            card = cards.nth(i)
            q_el = card.locator(".textq strong")
            if await q_el.count() > 0:
                q_text = (await q_el.inner_text()).strip()
                q_num = q_text.replace("Question No.", "").replace("Question no.", "").strip()
                if q_num == question_no or q_num == question_no.lstrip("0"):
                    target_card = card
                    break

        if target_card is None:
            if count == 1:
                target_card = cards.first
            else:
                return f"❌ Could not find Q{question_no} among {count} cards"

        file_input = target_card.locator('input[type="file"]').first
        if await file_input.count() == 0:
            file_input = page.locator('input[type="file"]').first

        if await file_input.count() == 0:
            return "❌ No file input found on assignment page"

        await file_input.set_input_files(file_path)
        await page.wait_for_timeout(1000)

        for btn_sel in [
            ".btn-submit", 'button:has-text("SUBMIT")', 'button:has-text("Submit")',
            'button[type="submit"]',
        ]:
            btn = target_card.locator(btn_sel).first
            if await btn.count() > 0:
                try:
                    await btn.click(timeout=5000)
                    await page.wait_for_timeout(2500)
                    for indicator in [
                        "text=successfully", "text=submitted", "text=uploaded",
                        ".v-snack__content", ".alert-success",
                    ]:
                        if await page.locator(indicator).count() > 0:
                            return "✅ Assignment submitted successfully!"
                    return "✅ File uploaded and SUBMIT clicked — verify on VOLP."
                except Exception as exc:
                    return f"⚠️ File attached but submit failed: {exc}"

        return "⚠️ File attached but no SUBMIT button found — submit manually."


# ---------------------------------------------------------------------------
# AI Solver (OpenRouter + python-docx)
# ---------------------------------------------------------------------------


class VolpSolver:
    def __init__(self, api_key: str):
        self.api_key = api_key

    def solve(self, question: str, course: str, question_no: str) -> str:
        system = (
            f"You are a computer science student solving an assignment for the course: {course}. "
            "Write a clear, well-structured academic answer. "
            "Use proper headings, numbered steps, and code blocks where appropriate. "
            "Aim for a thorough answer that demonstrates understanding."
        )
        last_error: Optional[Exception] = None
        for model in OPENROUTER_FREE_MODELS:
            try:
                resp = requests.post(
                    OPENROUTER_URL,
                    headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
                    json={
                        "model": model,
                        "messages": [
                            {"role": "system", "content": system},
                            {"role": "user", "content": f"Question {question_no}:\n\n{question}"},
                        ],
                    },
                    timeout=120,
                )
                if resp.status_code != 200:
                    last_error = RuntimeError(f"Model {model} returned HTTP {resp.status_code}")
                    continue
                data = resp.json()
                if "error" in data:
                    last_error = RuntimeError(data["error"].get("message", str(data["error"])))
                    continue
                return data["choices"][0]["message"]["content"]
            except requests.RequestException as exc:
                last_error = exc
        raise last_error or RuntimeError("All models failed")

    def make_docx(
        self, question: str, answer: str, course: str, question_no: str, output_path: str
    ) -> None:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        title = doc.add_heading(f"{course}", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub = doc.add_paragraph(f"Assignment — Question {question_no}")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.size = Pt(12)

        doc.add_paragraph("")
        doc.add_heading("Question:", level=1)
        doc.add_paragraph(question)
        doc.add_paragraph("")
        doc.add_heading("Answer:", level=1)

        in_code = False
        code_lines: List[str] = []

        for line in answer.split("\n"):
            if line.startswith("```"):
                if in_code:
                    para = doc.add_paragraph("\n".join(code_lines))
                    para.style = "No Spacing"
                    for run in para.runs:
                        run.font.name = "Courier New"
                        run.font.size = Pt(9)
                    code_lines = []
                    in_code = False
                else:
                    in_code = True
            elif in_code:
                code_lines.append(line)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            else:
                doc.add_paragraph(line)

        if code_lines:
            doc.add_paragraph("\n".join(code_lines))

        doc.save(output_path)


# ---------------------------------------------------------------------------
# Shared Playwright helpers
# ---------------------------------------------------------------------------


async def _route_filter(route):
    if route.request.resource_type in {"image", "font", "media"}:
        await route.abort()
    else:
        await route.continue_()


async def _click_tab(page, tab_text: str) -> bool:
    for sel in [
        f'.v-tab:has-text("{tab_text}")',
        f'a:has-text("{tab_text}")',
        f'[role=tab]:has-text("{tab_text}")',
    ]:
        if await page.locator(sel).count() > 0:
            try:
                await page.locator(sel).first.click(timeout=5000)
                await page.wait_for_load_state("networkidle", timeout=30000)
                return True
            except Exception:
                pass
    return False


def _dedupe(assignments: List[Assignment]) -> List[Assignment]:
    seen: Dict[str, Assignment] = {}
    for item in assignments:
        seen[item.entity_key] = item
    return list(seen.values())


def _esc(text: str) -> str:
    for ch in r"\_*[]()~`>#+-=|{}.!":
        text = text.replace(ch, f"\\{ch}")
    return text


# ---------------------------------------------------------------------------
# Change detection + messaging
# ---------------------------------------------------------------------------


def detect_changes(
    old_rows: Dict[str, sqlite3.Row], new_items: List[Assignment]
) -> Dict[str, List[Assignment]]:
    created, updated, submission_changed = [], [], []
    for item in new_items:
        existing = old_rows.get(item.entity_key)
        if not existing:
            created.append(item)
            continue
        old_due = (existing["due_iso"] or "").strip() or (existing["due_raw"] or "").lower()
        new_due = (item.due_iso or "").strip() or item.due_raw.lower()
        if old_due != new_due:
            updated.append(item)
        old_file = (existing["submitted_file"] or "").strip()
        new_file = (item.submitted_file or "").strip()
        if new_file and new_file != old_file:
            submission_changed.append(item)
    return {"created": created, "updated": updated, "submission_changed": submission_changed}


def build_change_message(created: List[Assignment], updated: List[Assignment]) -> str:
    lines = ["📚 VOLP update detected"]
    if created:
        lines.append(f"\nNew assignments: {len(created)}")
        for item in created[:8]:
            due = item.due_raw or "no due date"
            lines.append(f"• {item.course} | Q{item.question_no} | Due: {due}")
    if updated:
        lines.append(f"\nUpdated deadlines: {len(updated)}")
        for item in updated[:8]:
            lines.append(f"• {item.course} | Q{item.question_no} | New due: {item.due_raw}")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Config + utilities
# ---------------------------------------------------------------------------


def load_config() -> dict:
    load_dotenv()
    return {
        "base_url": os.getenv("VOLP_BASE_URL", "https://classroom.volp.in"),
        "login_url": os.getenv("VOLP_LOGIN_URL", "https://classroom.volp.in/login"),
        "start_url": os.getenv("VOLP_START_URL", "https://classroom.volp.in/learner/my-courses"),
        "username": os.getenv("VOLP_USERNAME", ""),
        "password": os.getenv("VOLP_PASSWORD", ""),
        "telegram_token": os.getenv("TELEGRAM_BOT_TOKEN", ""),
        "telegram_chat_id": os.getenv("TELEGRAM_CHAT_ID", ""),
        "interval_minutes": int(os.getenv("CHECK_INTERVAL_MINUTES", "10")),
        "headless": os.getenv("HEADLESS", "true").lower() == "true",
        "db_path": os.getenv("STATE_DB_PATH", "./state.db"),
        "scan_max_courses": int(os.getenv("SCAN_MAX_COURSES", "20")),
        "scan_retries": int(os.getenv("SCAN_RETRIES", "2")),
        "volp_tz_offset_minutes": int(os.getenv("VOLP_TZ_OFFSET_MINUTES", "330")),
        "openrouter_key": os.getenv("OPENROUTER_API_KEY", ""),
    }


def normalize_due_to_iso(due_raw: str, tz_offset_minutes: int) -> str:
    value = (due_raw or "").strip()
    if not value:
        return ""
    parts = value.replace("  ", " ").strip().split()
    date_part = parts[0] if parts else ""
    time_part = parts[1] if len(parts) > 1 else "23:59"
    am_pm = parts[2].upper() if len(parts) > 2 else ""
    try:
        day, month, year = [int(x) for x in date_part.split("/")]
        hour, minute = [int(x) for x in time_part.split(":")]
    except Exception:
        return ""
    if am_pm == "PM" and hour != 12:
        hour += 12
    if am_pm == "AM" and hour == 12:
        hour = 0
    hour = hour % 24
    local = datetime(year, month, day, hour, minute, tzinfo=timezone(timedelta(minutes=tz_offset_minutes)))
    return local.astimezone(timezone.utc).isoformat()


# ---------------------------------------------------------------------------
# --once mode (single-user, backward compat)
# ---------------------------------------------------------------------------


async def run_once(config: dict, state: StateStore, telegram: TelegramClient) -> None:
    user_id = config.get("telegram_chat_id") or ""
    scanner = VolpScanner(
        base_url=config["base_url"],
        login_url=config["login_url"],
        start_url=config["start_url"],
        username=config["username"],
        password=config["password"],
        headless=config["headless"],
        max_courses=max(1, config["scan_max_courses"]),
        tz_offset_minutes=config["volp_tz_offset_minutes"],
    )

    print("[run] scanning VOLP...")
    new_items: Optional[List[Assignment]] = None
    for attempt in range(1, config["scan_retries"] + 1):
        try:
            new_items = await scanner.scan()
            break
        except Exception as exc:
            print(f"[run] attempt {attempt} failed: {exc}")
            await asyncio.sleep(attempt * 2)

    if new_items is None:
        raise RuntimeError("All scan attempts failed")

    print(f"[run] found {len(new_items)} total assignment records")
    for a in new_items:
        due = a.due_raw or "no due date"
        print(f"  [{a.course}] Q{a.question_no} — {a.title[:60]} — Due: {due}")

    old_rows = state.load_assignments(user_id)
    changes = detect_changes(old_rows, new_items)
    state.upsert_assignments(new_items, user_id)

    if telegram.enabled() and state.get_meta("telegram_test_sent") != "1":
        ok = telegram.send("✅ VOLP Tracker connected.")
        if ok:
            state.set_meta("telegram_test_sent", "1")

    created = changes["created"]
    updated = changes["updated"]
    if created or updated:
        msg = build_change_message(created, updated)
        if telegram.enabled():
            telegram.send(msg)
        print(f"[run] changes: +{len(created)} new, ~{len(updated)} updated")
    else:
        print("[run] no changes detected")


# ---------------------------------------------------------------------------
# Telegram bot — /start registration ConversationHandler
# ---------------------------------------------------------------------------


async def _start_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Entry point: check if registered, else begin registration flow."""
    if update.message.chat.type != "private":
        await update.message.reply_text("Please start me in a private chat to register.")
        return ConversationHandler.END
    chat_id = str(update.effective_user.id)
    state: StateStore = context.bot_data["state"]
    user = state.get_user(chat_id)

    if user:
        status = user["status"]
        if status == "approved":
            await update.message.reply_text(
                "✅ You're already registered and approved!\n"
                "Use /status to see your assignments or /solve to list unsubmitted ones."
            )
        elif status == "pending":
            await update.message.reply_text(
                "⏳ Your registration is pending admin approval. Please wait."
            )
        elif status == "denied":
            await update.message.reply_text(
                "❌ Your registration was denied by the admin."
            )
        return ConversationHandler.END

    await update.message.reply_text(
        "👋 Welcome to VOLP Tracker!\n\n"
        "This bot tracks your VOLP assignments and lets you submit directly from Telegram.\n\n"
        "Please enter your VOLP email address:"
    )
    return WAIT_USERNAME


async def _handle_username(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Store email, ask for password."""
    context.user_data["volp_username"] = update.message.text.strip()
    await update.message.reply_text(
        "🔑 Now enter your VOLP password:\n"
        "_(Your message will be deleted immediately for security)_",
        parse_mode="Markdown",
    )
    return WAIT_PASSWORD


async def _handle_password(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Verify credentials, register user, notify admin."""
    password = update.message.text.strip()
    chat_id = str(update.effective_user.id)
    username = context.user_data.get("volp_username", "")
    config: dict = context.bot_data["config"]
    state: StateStore = context.bot_data["state"]

    # Delete the password message immediately
    try:
        await update.message.delete()
    except Exception:
        pass

    verifying_msg = await update.effective_chat.send_message("⏳ Verifying your VOLP credentials…")

    try:
        valid = await _verify_volp_credentials(username, password, config)
    except Exception as exc:
        await verifying_msg.edit_text(f"❌ Verification error: {exc}\n\nSend /start to try again.")
        context.user_data.clear()
        return ConversationHandler.END

    if not valid:
        await verifying_msg.edit_text(
            "❌ Login failed — please check your VOLP email and password.\n\n"
            "Send /start to try again."
        )
        context.user_data.clear()
        return ConversationHandler.END

    # Register as pending
    state.register_user(chat_id, username, password, status="pending")

    # Notify admin
    admin_chat_id = config["telegram_chat_id"]
    tg_user = update.effective_user
    display_name = tg_user.full_name or tg_user.username or f"ID:{chat_id}"
    kb = InlineKeyboardMarkup([[
        InlineKeyboardButton("✅ Approve", callback_data=f"approve:{chat_id}"),
        InlineKeyboardButton("❌ Deny", callback_data=f"deny:{chat_id}"),
    ]])
    try:
        await context.bot.send_message(
            chat_id=admin_chat_id,
            text=(
                f"🆕 New registration request:\n"
                f"👤 {display_name} (Telegram ID: {chat_id})\n"
                f"📧 VOLP: {username}\n\n"
                "Approve or deny this user?"
            ),
            reply_markup=kb,
        )
    except Exception as exc:
        print(f"[registration] failed to notify admin {admin_chat_id}: {exc}")

    await verifying_msg.edit_text(
        "✅ Credentials verified! Your request has been sent to the admin.\n"
        "You'll receive a message once you're approved."
    )
    context.user_data.clear()
    return ConversationHandler.END


async def _cancel_registration(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    context.user_data.clear()
    await update.message.reply_text("Registration cancelled.")
    return ConversationHandler.END


# ---------------------------------------------------------------------------
# Admin approve / deny callbacks
# ---------------------------------------------------------------------------


async def _handle_approve(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    await query.answer()

    admin_id = str(query.from_user.id)
    config: dict = context.bot_data["config"]
    admin_cfg = str(config["telegram_chat_id"])
    if not admin_cfg or admin_id != admin_cfg:
        await query.answer("Unauthorized.", show_alert=True)
        return

    user_chat_id = query.data.split(":", 1)[1]
    state: StateStore = context.bot_data["state"]
    state.update_user_status(user_chat_id, "approved")

    await query.edit_message_text(query.message.text + "\n\n✅ Approved!")

    try:
        await context.bot.send_message(
            chat_id=user_chat_id,
            text=(
                "✅ Your VOLP Tracker account has been approved!\n"
                "Your first scan will run shortly.\n\n"
                "Commands:\n"
                "/status — view all assignments\n"
                "/solve — list unsubmitted assignments\n"
                "/setkey <key> — set your OpenRouter API key for AI Solve"
            ),
        )
    except Exception as exc:
        print(f"[approve] failed to notify user {user_chat_id}: {exc}")

    # Schedule silent first scan for this user
    context.job_queue.run_once(
        _scan_user_job,
        when=5,
        data={"user_chat_id": user_chat_id, "silent": True},
        name=f"first_scan_{user_chat_id}",
    )


async def _handle_deny(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    await query.answer()

    admin_id = str(query.from_user.id)
    config: dict = context.bot_data["config"]
    admin_cfg = str(config["telegram_chat_id"])
    if not admin_cfg or admin_id != admin_cfg:
        await query.answer("Unauthorized.", show_alert=True)
        return

    user_chat_id = query.data.split(":", 1)[1]
    state: StateStore = context.bot_data["state"]
    state.update_user_status(user_chat_id, "denied")

    await query.edit_message_text(query.message.text + "\n\n❌ Denied.")

    try:
        await context.bot.send_message(
            chat_id=user_chat_id,
            text="❌ Your VOLP Tracker registration was denied by the admin.",
        )
    except Exception as exc:
        print(f"[deny] failed to notify user {user_chat_id}: {exc}")


# ---------------------------------------------------------------------------
# Scan jobs — per-user
# ---------------------------------------------------------------------------


async def _scan_job(context: ContextTypes.DEFAULT_TYPE) -> None:
    """Periodic scan for all approved users."""
    config: dict = context.bot_data["config"]
    state: StateStore = context.bot_data["state"]

    users = state.get_all_users(status="approved")
    print(f"[scan_job] scanning {len(users)} approved user(s)")

    for user in users:
        try:
            await _scan_user(context, user, config, silent=False)
        except Exception as exc:
            print(f"[scan_job] user {user['chat_id']} failed: {exc}")


async def _scan_user_job(context: ContextTypes.DEFAULT_TYPE) -> None:
    """One-shot scan for a specific user (used after approval)."""
    data = context.job.data
    user_chat_id = data["user_chat_id"]
    silent = data.get("silent", False)

    config: dict = context.bot_data["config"]
    state: StateStore = context.bot_data["state"]

    user = state.get_user(user_chat_id)
    if not user or user["status"] != "approved":
        return

    try:
        await _scan_user(context, user, config, silent=silent)
    except Exception as exc:
        print(f"[scan_user_job] user {user_chat_id} failed: {exc}")


async def _scan_user(context, user, config: dict, silent: bool = False) -> None:
    """Scan VOLP for one user and send Telegram notifications."""
    chat_id = user["chat_id"]
    state: StateStore = context.bot_data["state"]

    # Per-user scan lock prevents concurrent scans for the same account
    locks: dict = context.bot_data.setdefault("scan_locks", {})
    if chat_id not in locks:
        locks[chat_id] = asyncio.Lock()
    lock = locks[chat_id]
    if lock.locked():
        print(f"[scan_user] scan already running for {chat_id}, skipping")
        return

    async with lock:
        scanner = VolpScanner(
            base_url=config["base_url"],
            login_url=config["login_url"],
            start_url=config["start_url"],
            username=user["volp_username"],
            password=user["volp_password"],
            headless=config["headless"],
            max_courses=max(1, config["scan_max_courses"]),
            tz_offset_minutes=config["volp_tz_offset_minutes"],
        )

        print(f"[scan_user] scanning for {chat_id} ({user['volp_username']})...")
        try:
            new_items = await scanner.scan()
        except Exception as exc:
            print(f"[scan_user] {chat_id} scan failed: {exc}")
            return

        print(f"[scan_user] {chat_id}: {len(new_items)} records")
        old_rows = state.load_assignments(user_id=chat_id)
        changes = detect_changes(old_rows, new_items)
        state.upsert_assignments(new_items, user_id=chat_id)
        state.update_user_last_scan(chat_id)

        if silent:
            print(f"[scan_user] {chat_id}: silent scan done (no notifications)")
            return

        # Send "connected" message once per user
        meta_key = f"connected_sent_{chat_id}"
        if state.get_meta(meta_key) != "1":
            await context.bot.send_message(
                chat_id=chat_id,
                text="✅ VOLP Tracker connected. I'll notify you when assignments change.",
            )
            state.set_meta(meta_key, "1")

        def assignment_kb(key: str) -> InlineKeyboardMarkup:
            return InlineKeyboardMarkup([[
                InlineKeyboardButton("📤 Submit", callback_data=f"submit:{key}"),
                InlineKeyboardButton("🤖 AI Solve", callback_data=f"aisolve:{key}"),
            ]])

        created = changes["created"]
        updated = changes["updated"]

        # Rate throttle: send a single summary when there are many changes to
        # avoid hitting Telegram's 429 flood limit on first scan
        if len(created) + len(updated) > 8:
            msg = build_change_message(created, updated)
            await context.bot.send_message(chat_id=chat_id, text=msg)
        else:
            for item in created:
                due = item.due_raw or "no due date"
                text = (
                    f"🆕 *New assignment\\!*\n"
                    f"📚 {_esc(item.course)}\n"
                    f"❓ Q{_esc(item.question_no)}: {_esc(item.title[:80])}\n"
                    f"⏰ Due: {_esc(due)}"
                )
                await context.bot.send_message(
                    chat_id=chat_id, text=text,
                    parse_mode="MarkdownV2", reply_markup=assignment_kb(item.entity_key),
                )
                await asyncio.sleep(0.05)

            for item in updated:
                text = (
                    f"📝 *Deadline updated*\n"
                    f"📚 {_esc(item.course)}\n"
                    f"❓ Q{_esc(item.question_no)}\n"
                    f"⏰ New due: {_esc(item.due_raw)}"
                )
                await context.bot.send_message(
                    chat_id=chat_id, text=text,
                    parse_mode="MarkdownV2", reply_markup=assignment_kb(item.entity_key),
                )
                await asyncio.sleep(0.05)

        for item in changes.get("submission_changed", []):
            text = (
                f"✅ *Submission recorded\\!*\n"
                f"📚 {_esc(item.course)}\n"
                f"❓ Q{_esc(item.question_no)}: {_esc(item.title[:60])}\n"
                f"📎 File: {_esc(item.submitted_file)}"
            )
            await context.bot.send_message(chat_id=chat_id, text=text, parse_mode="MarkdownV2")

        if not created and not updated and not changes.get("submission_changed"):
            print(f"[scan_user] {chat_id}: no changes")


# ---------------------------------------------------------------------------
# Telegram bot handlers — all scoped by chat_id + approval check
# ---------------------------------------------------------------------------


def _require_approved(state: StateStore, chat_id: str) -> Optional[sqlite3.Row]:
    """Return user row if approved, else None."""
    user = state.get_user(chat_id)
    if user and user["status"] == "approved":
        return user
    return None


async def _handle_submit_click(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    await query.answer()

    chat_id = str(query.message.chat_id)
    state: StateStore = context.bot_data["state"]
    user = _require_approved(state, chat_id)
    if not user:
        await query.answer("You need to be approved to use this.", show_alert=True)
        return

    entity_key = query.data.split(":", 1)[1]
    row = state.get_by_entity_key(entity_key, chat_id)
    if not row:
        await query.message.reply_text("❌ Assignment not found in DB.")
        return

    context.user_data["pending_submit"] = entity_key
    await query.message.reply_text(
        f"📎 Send me the file for:\n"
        f"*{_esc(row['course'])}* — Q{_esc(row['question_no'])}\n"
        f"Due: {_esc(row['due_raw'] or 'no due date')}\n\n"
        "Attach as a *file* \\(not a photo\\)\\.",
        parse_mode="MarkdownV2",
    )


async def _handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = str(update.message.chat_id)
    state: StateStore = context.bot_data["state"]
    config: dict = context.bot_data["config"]

    user = _require_approved(state, chat_id)
    if not user:
        return

    entity_key = context.user_data.get("pending_submit")
    if not entity_key:
        await update.message.reply_text(
            "No pending submission\\. Click '📤 Submit' on an assignment first\\.",
            parse_mode="MarkdownV2",
        )
        return

    row = state.get_by_entity_key(entity_key, chat_id)
    if not row:
        await update.message.reply_text("❌ Assignment not found.")
        context.user_data.pop("pending_submit", None)
        return

    doc = update.message.document
    tg_file = await context.bot.get_file(doc.file_id)
    suffix = f"_{doc.file_name}" if doc.file_name else ".bin"
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    tmp.close()
    await tg_file.download_to_drive(tmp.name)

    await update.message.reply_text(
        f"⏳ Submitting *{_esc(doc.file_name or 'file')}* to VOLP\\.\\.\\.\n"
        f"{_esc(row['course'])} — Q{_esc(row['question_no'])}",
        parse_mode="MarkdownV2",
    )

    submitter = VolpSubmitter(
        login_url=config["login_url"],
        start_url=config["start_url"],
        username=user["volp_username"],
        password=user["volp_password"],
        headless=config["headless"],
    )

    try:
        result = await submitter.submit_file(
            overview_url=row["overview_url"] or "",
            source_url=row["source_url"] or "",
            question_no=row["question_no"] or "",
            file_path=tmp.name,
        )
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass

    context.user_data.pop("pending_submit", None)
    await update.message.reply_text(result)


async def _handle_ai_solve(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    query = update.callback_query
    await query.answer()

    chat_id = str(query.message.chat_id)
    state: StateStore = context.bot_data["state"]
    config: dict = context.bot_data["config"]

    user = _require_approved(state, chat_id)
    if not user:
        await query.answer("You need to be approved to use this.", show_alert=True)
        return

    entity_key = query.data.split(":", 1)[1]
    row = state.get_by_entity_key(entity_key, chat_id)
    if not row:
        await query.message.reply_text("❌ Assignment not found in DB.")
        return

    # Per-user key falls back to global env key
    api_key = (user["openrouter_key"] or "").strip() or config.get("openrouter_key", "")
    if not api_key:
        await query.message.reply_text(
            "❌ No OpenRouter API key set\\.\n"
            "Run: /setkey sk\\-or\\-\\.\\.\\.",
            parse_mode="MarkdownV2",
        )
        return

    question = (row["full_text"] or row["title"] or "").strip()
    if not question:
        await query.message.reply_text("❌ No problem statement found for this assignment.")
        return

    await query.message.reply_text(
        f"🤖 Generating AI answer for *{_esc(row['course'])}* Q{_esc(row['question_no'])}\\.\\.\\.",
        parse_mode="MarkdownV2",
    )

    solver = VolpSolver(api_key)
    try:
        answer = await asyncio.get_event_loop().run_in_executor(
            None, solver.solve, question, row["course"], row["question_no"]
        )
    except Exception as exc:
        await query.message.reply_text(f"❌ AI call failed: {exc}")
        return

    suffix = f"_Q{row['question_no']}.docx"
    tmp = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    tmp.close()
    try:
        solver.make_docx(question, answer, row["course"], row["question_no"], tmp.name)
    except Exception as exc:
        os.unlink(tmp.name)
        await query.message.reply_text(f"❌ Document generation failed: {exc}")
        return

    await query.message.reply_text("📄 Document ready\\. Submitting to VOLP\\.\\.\\.", parse_mode="MarkdownV2")

    submitter = VolpSubmitter(
        login_url=config["login_url"],
        start_url=config["start_url"],
        username=user["volp_username"],
        password=user["volp_password"],
        headless=config["headless"],
    )
    try:
        result = await submitter.submit_file(
            overview_url=row["overview_url"] or "",
            source_url=row["source_url"] or "",
            question_no=row["question_no"] or "",
            file_path=tmp.name,
        )
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass

    await query.message.reply_text(result)


async def _handle_setkey(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = str(update.message.chat_id)
    state: StateStore = context.bot_data["state"]

    user = _require_approved(state, chat_id)
    if not user:
        await update.message.reply_text(
            "❌ You need to be registered and approved first. Send /start to register."
        )
        return

    if not context.args:
        await update.message.reply_text("Usage: /setkey <openrouter_api_key>")
        return

    key = context.args[0].strip()
    state.set_user_key(chat_id, key)
    await update.message.reply_text("✅ OpenRouter API key saved! You can now use 🤖 AI Solve.")


async def _handle_solve(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Show unsubmitted assignments with AI Solve buttons."""
    chat_id = str(update.message.chat_id)
    state: StateStore = context.bot_data["state"]

    user = _require_approved(state, chat_id)
    if not user:
        await update.message.reply_text(
            "❌ You need to be registered and approved first. Send /start to register."
        )
        return

    rows = state.conn.execute(
        "SELECT entity_key, course, question_no, title, due_raw, submitted_file "
        "FROM assignments WHERE (submitted_file = '' OR submitted_file IS NULL) AND user_id = ? "
        "ORDER BY course, CAST(question_no AS INTEGER)",
        (chat_id,),
    ).fetchall()

    if not rows:
        await update.message.reply_text("🎉 All assignments have submissions!")
        return

    await update.message.reply_text(
        f"❌ <b>{len(rows)} unsubmitted assignment(s):</b>",
        parse_mode="HTML",
    )

    for row in rows:
        due = row["due_raw"] or "no due date"
        text = f"📚 <b>{row['course']}</b>\n❓ Q{row['question_no']}: {(row['title'] or '')[:80]}\n⏰ {due}"
        kb = InlineKeyboardMarkup([[
            InlineKeyboardButton("📤 Submit", callback_data=f"submit:{row['entity_key']}"),
            InlineKeyboardButton("🤖 AI Solve", callback_data=f"aisolve:{row['entity_key']}"),
        ]])
        await update.message.reply_text(text, parse_mode="HTML", reply_markup=kb)


async def _handle_status(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    chat_id = str(update.message.chat_id)
    state: StateStore = context.bot_data["state"]

    user = _require_approved(state, chat_id)
    if not user:
        await update.message.reply_text(
            "❌ You need to be registered and approved first. Send /start to register."
        )
        return

    rows = state.conn.execute(
        "SELECT course, question_no, title, due_raw, submitted_file "
        "FROM assignments WHERE user_id = ? "
        "ORDER BY course, CAST(question_no AS INTEGER)",
        (chat_id,),
    ).fetchall()

    if not rows:
        await update.message.reply_text("No assignments tracked yet. Wait for the next scan.")
        return

    by_course: dict = defaultdict(list)
    for row in rows:
        by_course[row["course"]].append(row)

    submitted = sum(1 for r in rows if r["submitted_file"])
    total = len(rows)

    lines = [f"📊 <b>Assignment Status</b>  ✅ {submitted}/{total} submitted\n"]
    for course, assignments in by_course.items():
        lines.append(f"\n<b>{course}</b>")
        for a in assignments:
            icon = "✅" if a["submitted_file"] else "❌"
            due = a["due_raw"] or "no due date"
            fname = f"  📎 {a['submitted_file']}" if a["submitted_file"] else ""
            lines.append(f"  {icon} Q{a['question_no']} | {due}{fname}")

    msg = "\n".join(lines)
    while msg:
        chunk, msg = msg[:3900], msg[3900:]
        await update.message.reply_text(chunk, parse_mode="HTML")


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def main() -> int:
    parser = argparse.ArgumentParser(description="VOLP assignment tracker + submission bot")
    parser.add_argument("--once", action="store_true", help="One scan then exit")
    args = parser.parse_args()

    config = load_config()
    state = StateStore(config["db_path"])

    if not config["telegram_token"]:
        print("[config] TELEGRAM_BOT_TOKEN is required")
        return 1

    # Backward compat: auto-register admin from .env if VOLP credentials provided
    admin_chat_id = config["telegram_chat_id"]
    if admin_chat_id and config["username"] and config["password"]:
        existing = state.get_user(admin_chat_id)
        if not existing:
            print(f"[startup] auto-registering admin {admin_chat_id} from env credentials")
            state.register_user(
                admin_chat_id, config["username"], config["password"], status="approved"
            )
            state.migrate_legacy_user(admin_chat_id)
            print(f"[startup] admin registered and legacy assignments migrated")

    if args.once:
        if not config["username"] or not config["password"]:
            print("[config] VOLP_USERNAME and VOLP_PASSWORD required for --once mode")
            return 1
        telegram = TelegramClient(config["telegram_token"], admin_chat_id)
        try:
            asyncio.run(run_once(config, state, telegram))
        except KeyboardInterrupt:
            pass
        except Exception as exc:
            print(f"[run] fatal: {exc}")
            return 1
        return 0

    if not admin_chat_id:
        print("[config] TELEGRAM_CHAT_ID required for daemon mode (identifies admin)")
        return 1

    app = ApplicationBuilder().token(config["telegram_token"]).build()
    app.bot_data["config"] = config
    app.bot_data["state"] = state
    app.bot_data["scan_locks"] = {}

    # Registration ConversationHandler (must be first)
    conv_handler = ConversationHandler(
        entry_points=[CommandHandler("start", _start_cmd)],
        states={
            WAIT_USERNAME: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, _handle_username),
                CommandHandler("cancel", _cancel_registration),
            ],
            WAIT_PASSWORD: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, _handle_password),
                CommandHandler("cancel", _cancel_registration),
            ],
        },
        fallbacks=[CommandHandler("cancel", _cancel_registration)],
        conversation_timeout=600,
    )
    app.add_handler(conv_handler)

    # Admin approve/deny buttons
    app.add_handler(CallbackQueryHandler(_handle_approve, pattern="^approve:"))
    app.add_handler(CallbackQueryHandler(_handle_deny, pattern="^deny:"))

    # User commands
    app.add_handler(CommandHandler("status", _handle_status))
    app.add_handler(CommandHandler("solve", _handle_solve))
    app.add_handler(CommandHandler("setkey", _handle_setkey))

    # Inline action buttons
    app.add_handler(CallbackQueryHandler(_handle_submit_click, pattern="^submit:"))
    app.add_handler(CallbackQueryHandler(_handle_ai_solve, pattern="^aisolve:"))

    # File upload handler
    app.add_handler(MessageHandler(filters.Document.ALL, _handle_document))

    # Periodic scan
    app.job_queue.run_repeating(
        _scan_job,
        interval=max(60, config["interval_minutes"] * 60),
        first=15,
    )

    print(f"[bot] starting (scan every {config['interval_minutes']} min, admin={admin_chat_id})")
    app.run_polling(allowed_updates=Update.ALL_TYPES)
    return 0


if __name__ == "__main__":
    sys.exit(main())
